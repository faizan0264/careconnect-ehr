import sys
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
        run.font.size = Pt(15)
        run.font.bold = True
    elif level == 2:
        run.font.color.rgb = RGBColor(27, 85, 155) # Medical Blue
        run.font.size = Pt(12)
        run.font.bold = True
    elif level == 3:
        run.font.color.rgb = RGBColor(41, 128, 185) # Accent
        run.font.size = Pt(10.5)
        run.font.bold = True
    return h

def create_table_styled(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79") # Deep Blue
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    # Data Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F4F7FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(40, 40, 40)

    # Column Widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EBF3FB")
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="1B559B"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"{title}\n")
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(27, 85, 155)
    run_title.font.size = Pt(9.5)
    
    run_body = p.add_run(text)
    run_body.font.size = Pt(9)
    run_body.font.color.rgb = RGBColor(40, 40, 40)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_submission_document(output_path):
    doc = Document()

    # Configure Margins (0.75 in for neat executive look)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_sub = title_p.add_run("PROJECT SPECIFICATION & EVALUATION BASELINE\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(27, 85, 155)

    run_title = title_p.add_run("CareConnect: Patient-Provider EHR")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_before = Pt(2)
    desc_p.paragraph_format.space_after = Pt(10)
    run_desc = desc_p.add_run("Official System Modules, Use Cases, and Evaluation Traceability Matrix")
    run_desc.font.size = Pt(11)
    run_desc.font.italic = True
    run_desc.font.color.rgb = RGBColor(90, 90, 90)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(8)
    run_line = p_div.add_run("―" * 58)
    run_line.font.color.rgb = RGBColor(200, 210, 220)

    # Metadata Submission Card
    create_table_styled(doc, ["Submission Parameter", "Details"], [
        ["Project Title", "CareConnect: Web-Based Patient-Provider Electronic Health Record (EHR)"],
        ["Target Industry Domain", "Healthcare & Life Sciences (Clinical Care & Patient Engagement)"],
        ["Technology Stack", "Frontend: Angular 17+ | Backend: Spring Boot 3.x (Java 17/21) | Database: Oracle DB"],
        ["Core Architectural Scope", "Patient Records, Clinical Documentation (SOAP), CPOE, Medication Management, Patient Portal"],
        ["Compliance & Security", "Role-Based Access Control (RBAC), HIPAA-aligned PHI Audit Trail, Drug-Allergy Safety Check"]
    ], col_widths=[2.0, 4.5])

    add_callout(doc, "Evaluation Alignment Note", 
                "This document represents the definitive scope baseline for the CareConnect EHR project. "
                "Every module and use case specified below maps directly to functional components in the Angular frontend, "
                "REST controller endpoints in Spring Boot, and normalized schema tables in Oracle Database.")

    # ==========================================
    # SECTION 1: MASTER MODULE ARCHITECTURE
    # ==========================================
    add_heading_styled(doc, "1. Master System Modules Overview", level=1)
    doc.add_paragraph(
        "CareConnect is organized into six interconnected, non-redundant functional modules. "
        "This modular structure ensures clear separation of concerns between clinical staff workflows and patient self-service."
    )

    module_overview = [
        ["Module 1: Authentication & RBAC", "Provides secure identity verification, role-based route protection (Doctor, Nurse, Patient, Admin), and stateless JWT session management."],
        ["Module 2: Patient Records & Demographics", "Manages the Master Patient Index (MPI), automated Medical Record Number (MRN) generation, demographic profiles, and recorded allergy registries."],
        ["Module 3: Clinical Documentation (SOAP Notes)", "Enables encounter check-in, real-time physiological vitals capture with threshold alerts, and structured SOAP (Subjective, Objective, Assessment, Plan) documentation."],
        ["Module 4: CPOE (Physician Order Entry)", "Allows clinicians to place diagnostic laboratory, radiology, and procedural orders with priority levels (Routine, Urgent, STAT) and track test result fulfillment."],
        ["Module 5: Medication Management & e-Prescribing", "Provides digital prescription authoring with dosing schedules, duration calculations, active medication reconciliation, and real-time drug-allergy contraindication alerts."],
        ["Module 6: Patient Self-Service Portal", "Delivers an intuitive patient-facing portal to review active medications, view completed diagnostic lab reports, examine encounter summaries, and print after-visit instructions."]
    ]
    create_table_styled(doc, ["System Module", "Functional Scope & Responsibility"], module_overview, col_widths=[2.3, 4.2])

    # ==========================================
    # SECTION 2: CRISP USE CASES BY MODULE
    # ==========================================
    add_heading_styled(doc, "2. Detailed System Use Cases by Module", level=1)

    use_case_groups = [
        {
            "mod_num": "Module 1",
            "mod_name": "Authentication & Role-Based Access Control (RBAC)",
            "use_cases": [
                {
                    "id": "UC-1.1",
                    "title": "Role-Based User Authentication & Session Establishment",
                    "actor": "Doctor, Nurse, Patient, Administrator",
                    "objective": "Verify user credentials securely and issue a role-encoded session token.",
                    "inputs": "• Username (String)\n• Password (Plaintext, HTTPS)\n• User persona selection / Auto-detect",
                    "processing": "1. Spring Security authenticates username against BCrypt hash in APP_USERS.\n2. Verifies user role (ROLE_DOCTOR, ROLE_NURSE, ROLE_PATIENT, ROLE_ADMIN).\n3. Generates signed HMAC-SHA256 JWT containing userId, role, and claims.\n4. Frontend Angular AuthGuard stores token and configures dynamic navigation.",
                    "outputs": "• HTTP 200 OK + JWT Token + User Profile metadata.\n• UI redirects to role-appropriate workspace (/patients or /portal).\n• Top navigation bar displays logged-in user name and role badge.",
                    "acceptance": "System denies invalid credentials, issues valid JWT on success, and routes doctors to clinical workspace and patients to patient portal."
                },
                {
                    "id": "UC-1.2",
                    "title": "Role-Guarded Navigation & Protected Route Enforcement",
                    "actor": "Authenticated User (Attempting unauthorized privilege elevation)",
                    "objective": "Prevent unauthorized personas from accessing privileged clinical or administrative functions.",
                    "inputs": "Direct browser URL navigation (e.g., Patient attempting to navigate to /encounters or /admin).",
                    "processing": "1. Angular RoleGuard intercepts route activation.\n2. Compares active token role against required route roles array.\n3. Rejects activation if role authorization is insufficient.",
                    "outputs": "• Access denied redirect to /unauthorized.\n• Visual banner: 'Access Restricted - Insufficient clinical privileges.'",
                    "acceptance": "Patients cannot access provider chart views; nurses cannot sign doctor orders; unauthenticated users are kicked to /login."
                }
            ]
        },
        {
            "mod_num": "Module 2",
            "mod_name": "Patient Records & Demographics Management",
            "use_cases": [
                {
                    "id": "UC-2.1",
                    "title": "Patient Registration & Unique MRN Auto-Generation",
                    "actor": "Nurse, Front Desk Staff, Administrator",
                    "objective": "Enroll a new patient into the EHR and assign an immutable Medical Record Number (MRN).",
                    "inputs": "• First Name, Last Name (String, Required)\n• Date of Birth (Date, past only)\n• Gender (Male, Female, Other)\n• Blood Group (Dropdown: A+, B+, O+, AB+, etc.)\n• Phone Number (Masked input)\n• Recorded Allergies (Comma-separated text / tag chips)\n• Emergency Contact (Name and Phone)",
                    "processing": "1. Validates required demographic fields and DOB sanity.\n2. Backend sequence auto-generates unique MRN (e.g., 'MRN-2026-0042').\n3. Persists record in PATIENTS table.\n4. Asynchronously logs CREATE_PATIENT event in AUDIT_LOGS.",
                    "outputs": "• HTTP 201 Created + Persisted Patient Record + Generated MRN.\n• UI displays confirmation toast and auto-selects newly created patient chart.",
                    "acceptance": "Unique MRN is generated without collision, patient profile is saved in Oracle DB, and record immediately appears in search index."
                },
                {
                    "id": "UC-2.2",
                    "title": "Fast Patient Search & 360° Demographics Profile View",
                    "actor": "Doctor, Nurse",
                    "objective": "Quickly locate patient records by MRN or Name and view their longitudinal clinical summary.",
                    "inputs": "Search query string (MRN or Patient Name) with real-time debounced input.",
                    "processing": "1. Angular input debounces search query by 300ms.\n2. Queries GET /api/v1/patients?search={query}.\n3. Returns matching patient cards with high-visibility allergy tags.",
                    "outputs": "• Filtered patient list showing: MRN, Full Name, Age/Sex, Blood Group, and Allergies.\n• Selecting a patient locks the 'Patient 360° Header' banner across all clinical tabs.",
                    "acceptance": "Searching by partial name or full MRN returns correct patient within 500ms; selecting patient opens full clinical profile."
                }
            ]
        },
        {
            "mod_num": "Module 3",
            "mod_name": "Clinical Documentation & Vitals Management (SOAP Notes)",
            "use_cases": [
                {
                    "id": "UC-3.1",
                    "title": "Clinical Encounter Initiation & Triage Vitals Capture",
                    "actor": "Nurse, Triage Staff",
                    "objective": "Start a new patient encounter and document baseline physiological vital signs with visual threshold alerts.",
                    "inputs": "• Patient ID & Encounter Type (Outpatient, Emergency, Follow-up)\n• Chief Complaint / Reason for Visit\n• Vitals Input:\n   - Systolic / Diastolic Blood Pressure (mmHg)\n   - Heart Rate / Pulse (bpm)\n   - Body Temperature (°F)\n   - Oxygen Saturation (SpO2 %)\n   - Respiratory Rate (breaths/min)",
                    "processing": "1. Creates record in ENCOUNTERS with status 'IN_PROGRESS'.\n2. Validates vital numerical ranges.\n3. Evaluates abnormal vital thresholds (e.g. SpO2 < 92% or BP > 140/90).\n4. Persists baseline measurements in CLINICAL_NOTES.",
                    "outputs": "• Active Encounter session created.\n• Vitals Summary Card rendered with real-time color badges (Normal = Green, Warning = Amber, Critical = Red).",
                    "acceptance": "Encounter binds to selected patient, vitals are saved, and abnormal measurements clearly trigger visual warning badges."
                },
                {
                    "id": "UC-3.2",
                    "title": "Structured SOAP Note Clinical Documentation",
                    "actor": "Doctor (Attending Physician)",
                    "objective": "Record comprehensive clinical findings and treatment decisions using the standard SOAP framework.",
                    "inputs": "• Subjective (S): Patient history, reported symptoms, duration\n• Objective (O): Physical exam observations, clinical findings\n• Assessment (A): Physician diagnosis / Clinical impression\n• Plan (P): Treatment strategy, diagnostic tests, follow-up timeline",
                    "processing": "1. Validates that Assessment (A) is documented prior to saving.\n2. Upserts narrative CLOB fields into CLINICAL_NOTES for the encounter.\n3. Provides auto-draft saving every 5 seconds to prevent data loss.",
                    "outputs": "• Draft saved indicator with last updated timestamp.\n• Complete, formatted clinical narrative view integrated into the patient's legal chart.",
                    "acceptance": "Physician can document S, O, A, P fields cleanly, draft saves automatically, and persisted narrative updates database record."
                },
                {
                    "id": "UC-3.3",
                    "title": "Encounter Completion & Clinical Sign-Off",
                    "actor": "Doctor",
                    "objective": "Sign off on clinical documentation and formally close the encounter.",
                    "inputs": "'Sign & Complete Visit' action button + Final confirmation dialog.",
                    "processing": "1. Confirms all mandatory sections (Vitals, Assessment, Plan) are filled.\n2. Updates ENCOUNTERS status to 'COMPLETED'.\n3. Freezes CLINICAL_NOTES into read-only mode to preserve legal chart integrity.\n4. Writes DOCUMENT_SOAP event to AUDIT_LOGS.",
                    "outputs": "• Status pill transitions from 'IN_PROGRESS' (Blue) to 'COMPLETED' (Green).\n• Notes locked with electronic doctor sign-off stamp and completion timestamp.",
                    "acceptance": "Completed encounter cannot be edited; read-only summary is accessible to patient in portal."
                }
            ]
        },
        {
            "mod_num": "Module 4",
            "mod_name": "Computerized Physician Order Entry (CPOE)",
            "use_cases": [
                {
                    "id": "UC-4.1",
                    "title": "Diagnostic Order Placement with STAT Priority Flagging",
                    "actor": "Doctor (Physician)",
                    "objective": "Electronically order diagnostic laboratory tests, imaging studies, or clinical procedures directly within an encounter.",
                    "inputs": "• Order Category: Laboratory, Radiology, Procedure\n• Order Name: e.g., 'Complete Blood Count (CBC)', 'Chest X-Ray (AP/Lat)', 'Lipid Profile'\n• Priority Level: Routine, Urgent, STAT (Immediate)\n• Clinical Indication / Notes: Reason for ordering test",
                    "processing": "1. Verifies ordering provider has ROLE_DOCTOR.\n2. Creates record in CPOE_ORDERS with status 'PENDING'.\n3. High-priority 'STAT' orders trigger immediate emergency highlighting in the worklist.\n4. Logs CREATE_ORDER in AUDIT_LOGS.",
                    "outputs": "• HTTP 201 Created + Unique Order ID.\n• Order appears immediately in Encounter Active Orders list and Diagnostic Worklist.\n• STAT orders display with red pulsing alert badge.",
                    "acceptance": "Doctor can place single or multiple diagnostic orders, assign STAT/Routine priority, and see them listed under active encounter orders."
                },
                {
                    "id": "UC-4.2",
                    "title": "Order Fulfillment & Diagnostic Results Entry",
                    "actor": "Lab Technician, Nurse, Doctor",
                    "objective": "Document diagnostic test findings, record result values, and mark orders as fulfilled.",
                    "inputs": "• Order ID selection\n• Status update: 'IN_PROGRESS' or 'COMPLETED'\n• Result Value / Findings narrative (e.g. 'WBC: 11.2 K/uL, Hemoglobin: 14.5 g/dL. Normal limits.')\n• Result completion notes",
                    "processing": "1. Updates CPOE_ORDERS status to 'COMPLETED'.\n2. Stores result text/values in result_value CLOB.\n3. Sets completed_at timestamp.\n4. Automatically makes results available in patient chart and patient portal.",
                    "outputs": "• Order status changes to 'COMPLETED' (Green checkmark).\n• Diagnostic report viewable via 'View Results' modal in Doctor Chart & Patient Portal.",
                    "acceptance": "Entering test result transitions order to completed, timestamps fulfillment, and exposes findings to both doctor and patient."
                }
            ]
        },
        {
            "mod_num": "Module 5",
            "mod_name": "Medication Management & e-Prescribing",
            "use_cases": [
                {
                    "id": "UC-5.1",
                    "title": "Electronic Prescription Creation with Dosage & Duration",
                    "actor": "Doctor (Physician)",
                    "objective": "Generate structured digital prescriptions with specific clinical dosing instructions.",
                    "inputs": "• Drug Name: (e.g., 'Amoxicillin 500mg', 'Lisinopril 10mg', 'Metformin 500mg')\n• Dosage Strength: (e.g., '500 mg')\n• Frequency: (e.g., 'Once daily', 'Twice daily (BID)', 'Every 8 hours (TID)', 'PRN (As needed)')\n• Administration Route: (Oral, Intravenous, Inhalation, Topical)\n• Therapy Duration: Integer days (e.g., 7 days, 30 days)\n• Patient Instructions: (e.g., 'Take with food and full glass of water')",
                    "processing": "1. Validates all prescription parameters.\n2. Performs automated pre-check against patient allergy records.\n3. Persists record in PRESCRIPTIONS table with status 'ACTIVE'.\n4. Links prescription to current encounter and patient ID.",
                    "outputs": "• HTTP 201 Created + Active Prescription record.\n• Active Medication table updates dynamically with drug details, remaining days, and status badge.",
                    "acceptance": "Prescription is saved with full dosing parameters, marked as ACTIVE, and rendered in the patient's active drug profile."
                },
                {
                    "id": "UC-5.2",
                    "title": "Automated Drug-Allergy Safety Cross-Check & Clinical Override",
                    "actor": "Doctor (Physician)",
                    "objective": "Protect patient safety by detecting contraindications between prescribed drugs and recorded patient allergies.",
                    "inputs": "• Doctor submits prescription for a drug matching patient allergy (e.g., Patient has recorded allergy 'Penicillin'; Doctor prescribes 'Amoxicillin').\n• Physician Override Input: Override Reason selection + Legal confirmation checkbox.",
                    "processing": "1. Backend & Frontend cross-check drug name and pharmacological class against PATIENTS.allergies.\n2. If conflict detected: Intercepts prescription and displays High-Priority Red Warning Dialog.\n3. If physician overrides: Requires mandatory clinical justification and logs ALLERGY_OVERRIDE in AUDIT_LOGS.\n4. If physician cancels: Discards draft without saving.",
                    "outputs": "• UI Modal: 'CRITICAL ALLERGY ALERT: Patient allergic to Penicillin. Amoxicillin contraindication.'\n• Upon override: Prescription saved with prominent '⚠️ OVERRIDDEN ALLERGY' flag.",
                    "acceptance": "System successfully detects and alerts on drug-allergy conflict; blocks accidental saving; allows audited physician override."
                },
                {
                    "id": "UC-5.3",
                    "title": "Active Medication Reconciliation & Discontinuation",
                    "actor": "Doctor",
                    "objective": "Review active drug therapy and safely discontinue resolved or adverse medications.",
                    "inputs": "• Select active medication row\n• Click 'Discontinue Medication'\n• Mandatory reason: Adverse reaction, Therapy completed, Ineffective, Switched therapy",
                    "processing": "1. Updates PRESCRIPTIONS status from 'ACTIVE' to 'DISCONTINUED'.\n2. Records discontinuation reason and timestamp.\n3. Synchronizes updated medication status with Patient Portal.",
                    "outputs": "• Medication row moves to 'Discontinued / Past Meds' tab with strikethrough styling.\n• Success notification: 'Prescription discontinued.'",
                    "acceptance": "Discontinued medication immediately stops displaying as active and reflects as discontinued in the Patient Portal."
                }
            ]
        },
        {
            "mod_num": "Module 6",
            "mod_name": "Patient Self-Service Portal",
            "use_cases": [
                {
                    "id": "UC-6.1",
                    "title": "Patient Personal Health Dashboard & Care Team Summary",
                    "actor": "Patient (Self)",
                    "objective": "Provide patients with a clear, accessible overview of their personal health record.",
                    "inputs": "Patient signs in with patient credentials; landing route is /portal.",
                    "processing": "1. JWT claims resolve patient ID.\n2. Concurrently retrieves patient demographics, active medications count, recent completed lab results, and past encounter history.\n3. Formats clinical data into patient-friendly presentation cards.",
                    "outputs": "• Welcome greeting with Patient Name and MRN.\n• Metric cards: Active Prescriptions Count | Recent Completed Test Results | Primary Attending Physician.\n• Quick navigation tabs: [My Medications, My Lab Results, My Visits].",
                    "acceptance": "Patient views only their own health summary; data updates automatically when doctor saves notes or fulfills orders."
                },
                {
                    "id": "UC-6.2",
                    "title": "My Medications & Prescription Schedule View",
                    "actor": "Patient",
                    "objective": "Allow patients to view their current active medications with plain-language dosing schedules.",
                    "inputs": "Click 'My Medications' tab in Patient Portal.",
                    "processing": "Queries GET /api/v1/portal/dashboard to fetch prescriptions where status = 'ACTIVE'.",
                    "outputs": "• Visual medication schedule cards showing: Drug Name, Dosage, Instructions ('Take twice daily after meals'), and Days Remaining bar.",
                    "acceptance": "Patient sees active prescriptions clearly with dosage and instructions; discontinued meds are hidden from active list."
                },
                {
                    "id": "UC-6.3",
                    "title": "Diagnostic Lab Results Viewer & Clinical Summary Export",
                    "actor": "Patient",
                    "objective": "Enable patients to review completed lab/imaging reports and print an After-Visit Summary (AVS).",
                    "inputs": "Click 'My Lab Results' or 'Download Visit Summary' on a completed encounter.",
                    "processing": "1. Fetches completed CPOE order results and doctor's plan.\n2. Renders clean readable report modal.\n3. Formats print-optimized After-Visit Summary (AVS) via browser print.",
                    "outputs": "• Test result cards with completed date, test name, and doctor's remarks.\n• Clean printable PDF / print preview of visit summary.",
                    "acceptance": "Patient can inspect fulfilled lab results and generate a clean printable After-Visit Summary."
                }
            ]
        }
    ]

    for group in use_case_groups:
        add_heading_styled(doc, f"{group['mod_num']}: {group['mod_name']}", level=1)
        
        for uc in group["use_cases"]:
            add_heading_styled(doc, f"{uc['id']}: {uc['title']}", level=2)
            
            uc_data = [
                ["Primary Actor(s)", uc["actor"]],
                ["Clinical Objective", uc["objective"]],
                ["Input Parameters & Data", uc["inputs"]],
                ["Processing & Business Logic", uc["processing"]],
                ["Outputs & Visual Feedback", uc["outputs"]],
                ["Demonstrable Acceptance Criteria", uc["acceptance"]]
            ]
            create_table_styled(doc, ["Dimension", "Specification Details"], uc_data, col_widths=[1.8, 4.7])

    # ==========================================
    # SECTION 3: EVALUATION TRACEABILITY MATRIX
    # ==========================================
    add_heading_styled(doc, "3. Evaluation Traceability Matrix (Scope Baseline)", level=1)
    doc.add_paragraph(
        "To ensure seamless alignment during final project evaluation, the matrix below maps every module "
        "and use case to its corresponding Angular Frontend Component, Spring Boot Controller/Service, and Oracle DB Table."
    )

    matrix_headers = ["Use Case ID", "Module Name", "Angular UI Component", "Spring Boot REST Endpoint", "Oracle DB Table"]
    matrix_rows = [
        ["UC-1.1", "Auth & RBAC", "LoginComponent", "POST /api/v1/auth/login", "APP_USERS"],
        ["UC-1.2", "Auth & RBAC", "AuthGuard, RoleGuard", "JWT Bearer Verification", "APP_USERS"],
        ["UC-2.1", "Patient Records", "PatientRegistrationModal", "POST /api/v1/patients", "PATIENTS, AUDIT_LOGS"],
        ["UC-2.2", "Patient Records", "PatientListComponent, ChartHeader", "GET /api/v1/patients?search=", "PATIENTS"],
        ["UC-3.1", "Clinical Docs", "VitalsEntryCardComponent", "POST /api/v1/encounters", "ENCOUNTERS, CLINICAL_NOTES"],
        ["UC-3.2", "Clinical Docs", "SoapEditorTabsComponent", "PUT /api/v1/encounters/:id/notes", "CLINICAL_NOTES"],
        ["UC-3.3", "Clinical Docs", "EncounterDetailComponent", "POST /api/v1/encounters/:id/complete", "ENCOUNTERS, AUDIT_LOGS"],
        ["UC-4.1", "CPOE Orders", "CpoeOrderCartComponent", "POST /api/v1/orders", "CPOE_ORDERS, AUDIT_LOGS"],
        ["UC-4.2", "CPOE Orders", "OrderResultModalComponent", "PUT /api/v1/orders/:id/status", "CPOE_ORDERS"],
        ["UC-5.1", "Medications", "PrescribeModalComponent", "POST /api/v1/prescriptions", "PRESCRIPTIONS"],
        ["UC-5.2", "Medications", "AllergyWarningDialogComponent", "POST /api/v1/prescriptions (override)", "PRESCRIPTIONS, AUDIT_LOGS"],
        ["UC-5.3", "Medications", "MedicationListComponent", "PUT /api/v1/prescriptions/:id/discontinue", "PRESCRIPTIONS"],
        ["UC-6.1", "Patient Portal", "PatientDashboardComponent", "GET /api/v1/portal/dashboard", "PATIENTS, ENCOUNTERS"],
        ["UC-6.2", "Patient Portal", "MyMedsCardComponent", "GET /api/v1/portal/dashboard", "PRESCRIPTIONS"],
        ["UC-6.3", "Patient Portal", "MyLabResultsComponent", "GET /api/v1/portal/dashboard", "CPOE_ORDERS"]
    ]
    create_table_styled(doc, matrix_headers, matrix_rows, col_widths=[0.8, 1.2, 1.6, 1.7, 1.2])

    # Save Document
    doc.save(output_path)
    print(f"Official Submission Document successfully saved at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "CareConnect_Modules_and_UseCases_Submission.docx"
    build_submission_document(out)
