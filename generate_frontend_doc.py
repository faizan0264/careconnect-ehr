import sys
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
        run.font.size = Pt(16)
    elif level == 2:
        run.font.color.rgb = RGBColor(27, 85, 155) # Medical Blue
        run.font.size = Pt(13)
    elif level == 3:
        run.font.color.rgb = RGBColor(41, 128, 185) # Accent Blue
        run.font.size = Pt(11)
    return h

def create_table_styled(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79") # Deep Blue
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    # Data Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F4F7FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(40, 40, 40)

    # Column Widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EBF3FB")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="1B559B"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"UI Architecture Principle: {title}\n")
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(27, 85, 155)
    run_title.font.size = Pt(9.5)
    
    run_body = p.add_run(text)
    run_body.font.size = Pt(9)
    run_body.font.color.rgb = RGBColor(40, 40, 40)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_frontend_document(output_path):
    doc = Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_sub = title_p.add_run("CARECONNECT EHR • FRONTEND ARCHITECTURE & USE CASE SPECIFICATION\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(27, 85, 155)

    run_title = title_p.add_run("CareConnect Frontend Modules & UI Use Cases")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_before = Pt(2)
    desc_p.paragraph_format.space_after = Pt(10)
    run_desc = desc_p.add_run("Complete breakdown of Angular Frontend Modules, UI Components, User Interactions, Inputs, Validations, and Visual Outputs across the entire application lifecycle.")
    run_desc.font.size = Pt(10.5)
    run_desc.font.italic = True
    run_desc.font.color.rgb = RGBColor(90, 90, 90)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(10)
    run_line = p_div.add_run("―" * 58)
    run_line.font.color.rgb = RGBColor(200, 210, 220)

    # Modules Summary Matrix
    add_heading_styled(doc, "Frontend Module Architecture Map", level=1)
    
    summary_headers = ["Module Identifier", "Feature Name", "Primary Angular Route", "Target User Persona"]
    summary_rows = [
        ["MOD-01", "Authentication & Security Module", "/login, /unauthorized", "All Roles (Doctor, Nurse, Patient, Admin)"],
        ["MOD-02", "Patient Directory & Chart Header Module", "/patients, /patients/register", "Doctor, Nurse, Clinical Staff"],
        ["MOD-03", "Clinical Encounter & SOAP Documentation", "/encounters/:id, /encounters/new", "Doctor (SOAP), Nurse (Vitals & Triage)"],
        ["MOD-04", "CPOE Diagnostic Orders Module", "/encounters/:id/cpoe, /orders", "Doctor (Ordering), Nurse/Lab (Results)"],
        ["MOD-05", "Medication Management & e-Prescribing", "/encounters/:id/medications", "Doctor (Prescribing with Allergy Alert)"],
        ["MOD-06", "Patient Self-Service Portal Module", "/portal, /portal/meds, /portal/labs", "Patient Persona (Self-Access only)"],
        ["MOD-07", "HIPAA Audit & Administration Module", "/admin/audit, /admin/users", "System Administrator, Compliance Officer"]
    ]
    create_table_styled(doc, summary_headers, summary_rows, col_widths=[1.2, 2.3, 1.6, 1.8])

    add_callout(doc, "Angular Design Pattern", 
                "Each feature module follows the Smart (Container) / Dumb (Presentational) component pattern. "
                "Smart components interact with Angular Injectable Services and RxJS Observables, while Dumb components receive data via @Input() and emit events via @Output().")

    # ==========================================
    # MODULE DETAILS & USE CASES
    # ==========================================

    modules_data = [
        {
            "id": "MOD-01",
            "title": "Authentication & Session Security Module",
            "desc": "Handles secure persona identification, JWT token retention, role-based UI route guards, and automatic logout.",
            "components": "LoginComponent, UnauthorizedComponent, SessionTimeoutModalComponent, AuthGuard, RoleGuard, JwtInterceptor",
            "use_cases": [
                {
                    "code": "UC-FE-01.1",
                    "title": "Role-Aware User Login & Dynamic Navigation Redirect",
                    "actor": "Doctor, Nurse, Patient, Administrator",
                    "ui_inputs": "• Username (Text Input, required, minlength 3)\n• Password (Password Input with show/hide eye toggle, required)\n• Role Switcher / Auto-detect",
                    "ui_validation": "• Client-side dirty/touched error states\n• Inline error banner: 'Invalid username or password'\n• Disable 'Sign In' button while submission is in-flight (spinning loader)",
                    "ui_processing": "1. Angular Reactive Form captures credentials.\n2. AuthService invokes POST /api/v1/auth/login.\n3. On 200 OK: Store token and user metadata in localStorage/sessionStorage.\n4. Route resolution logic:\n   - If role == ROLE_DOCTOR or ROLE_NURSE -> Navigate to '/patients'\n   - If role == ROLE_PATIENT -> Navigate to '/portal'\n   - If role == ROLE_ADMIN -> Navigate to '/admin/audit'",
                    "ui_outputs": "• Smooth transition animation to designated dashboard\n• App Shell updates: Top navbar reflects user's Full Name, Role Badge, and Logout button"
                },
                {
                    "code": "UC-FE-01.2",
                    "title": "Unauthorized Route Access Prevention & Interception",
                    "actor": "Any logged-in user attempting URL tampering",
                    "ui_inputs": "Direct browser URL address bar manipulation (e.g., Patient typing '/admin/audit' or '/encounters/1001/cpoe')",
                    "ui_validation": "Angular RoleGuard inspects user role stored in AuthService before activating route",
                    "ui_processing": "1. RoleGuard verifies requested route's 'data: { roles: [...] }' against current user role.\n2. If mismatch detected, immediately aborts route activation.",
                    "ui_outputs": "• Redirect to '/unauthorized' screen\n• Visual Display: 'Access Restricted - You do not have clinical privileges to view this page.'\n• 'Return to Dashboard' primary button"
                },
                {
                    "code": "UC-FE-01.3",
                    "title": "Session Expiry & Token Injection Interceptor",
                    "actor": "Active User across any screen",
                    "ui_inputs": "Any outgoing HTTP request triggered by user interaction",
                    "ui_validation": "Token expiration check (decoded JWT exp claim)",
                    "ui_processing": "1. JwtInterceptor clones every outgoing request and injects 'Authorization: Bearer <token>'.\n2. If server returns HTTP 401 Unauthorized or 403 Forbidden, interceptor catches error, purges local session, and emits logout event.",
                    "ui_outputs": "• Toast notification: 'Your session has expired. Please sign in again.'\n• Immediate redirect to '/login'"
                }
            ]
        },
        {
            "id": "MOD-02",
            "title": "Patient Directory & Master Chart Header Module",
            "desc": "Provides rapid patient discovery by name/MRN, new patient registration, and a sticky clinical header visible throughout encounters.",
            "components": "PatientListComponent, PatientSearchInputComponent, PatientRegistrationModalComponent, PatientChartHeaderComponent",
            "use_cases": [
                {
                    "code": "UC-FE-02.1",
                    "title": "Real-Time Patient Search & Directory Filtering",
                    "actor": "Doctor, Nurse",
                    "ui_inputs": "• Search input field with debounce (300ms)\n• Filter chips: 'All', 'Admitted', 'Outpatient', 'High Risk Allergies'\n• Pagination controls (10/25/50 per page)",
                    "ui_validation": "Sanitize query string; show minimum 2 characters hint",
                    "ui_processing": "1. Search box binds via RxJS 'debounceTime(300)' and 'distinctUntilChanged()'.\n2. PatientService triggers GET /api/v1/patients?search={query}.\n3. Table displays skeletal loader while data fetches.",
                    "ui_outputs": "• Reactive table displaying: MRN, Full Name, Age/Gender, Blood Group, Recorded Allergies Badge, and 'Open Chart' action button\n• Zero results display: 'No patients found matching query. Click + New Patient to register.'"
                },
                {
                    "code": "UC-FE-02.2",
                    "title": "New Patient Registration Form & Auto-Generated MRN",
                    "actor": "Nurse, Front Desk Staff",
                    "ui_inputs": "• First Name & Last Name (Text, required)\n• Date of Birth (Date picker, cannot be in future)\n• Gender (Radio: Male, Female, Other)\n• Blood Group (Dropdown: A+, A-, B+, B-, O+, O-, AB+, AB-)\n• Primary Phone (Masked input: (999) 999-9999)\n• Allergies (Tag input with autocomplete, e.g. 'Penicillin', 'Sulfa', 'Latex')\n• Emergency Contact (Name and Phone)",
                    "ui_validation": "• Strict reactive validation on all required fields\n• Red border and error message on invalid blur\n• Submit button disabled until form valid",
                    "ui_processing": "1. Submits payload to POST /api/v1/patients.\n2. On success, closes modal, adds new patient to top of list, and emits toast.",
                    "ui_outputs": "• Success Toast: 'Patient registered successfully with MRN-2026-0042'\n• Direct navigation option: 'Open New Encounter Now'"
                },
                {
                    "code": "UC-FE-02.3",
                    "title": "Persistent 360° Patient Chart Header Display",
                    "actor": "Doctor, Nurse (Context across Encounters, CPOE, Meds)",
                    "ui_inputs": "Active patient selection from any clinical workflow",
                    "ui_validation": "Patient ID present in route param or active state store",
                    "ui_processing": "Loads patient demographic summary once per chart session and pins to top of screen.",
                    "ui_outputs": "• Sticky top bar showing: Patient Name, MRN, Age/DOB, Sex, Blood Group, Emergency Contact\n• Prominent Pulsing Red/Amber Allergy Badge (e.g., '⚠️ ALLERGIES: Penicillin, NSAIDs')\n• Quick action: 'End Encounter' button"
                }
            ]
        },
        {
            "id": "MOD-03",
            "title": "Clinical Encounter & SOAP Documentation Module",
            "desc": "The primary clinical workspace where nurses enter vitals and physicians record structured Subjective, Objective, Assessment, and Plan notes.",
            "components": "EncounterDetailComponent, VitalsEntryCardComponent, SoapEditorTabsComponent, EncounterHistoryTimelineComponent",
            "use_cases": [
                {
                    "code": "UC-FE-03.1",
                    "title": "Clinical Triage & Real-Time Vitals Input Grid",
                    "actor": "Nurse, Triage Specialist",
                    "ui_inputs": "• Systolic BP (Number: 40-300 mmHg)\n• Diastolic BP (Number: 20-200 mmHg)\n• Pulse / Heart Rate (Number: 30-250 bpm)\n• Temperature (Number: 90.0-108.0 °F)\n• SpO2 (Number: 50-100 %)\n• Respiratory Rate (Number: 8-50 breaths/min)",
                    "ui_validation": "• Real-time range validation with color indicators:\n  - Normal: Green border\n  - High/Low Alert: Amber border\n  - Critical Alert: Red pulsing border (e.g. SpO2 < 90% or Systolic > 180)",
                    "ui_processing": "1. Computes mean arterial pressure (MAP) automatically on client.\n2. Auto-saves draft vitals to local component state on change.",
                    "ui_outputs": "• Vitals stat card updates in real-time\n• Abnormal value flags displayed: 'High Blood Pressure (Stage 2 Hypertension)'"
                },
                {
                    "code": "UC-FE-03.2",
                    "title": "Structured SOAP Note Authoring & Clinical Narrative",
                    "actor": "Doctor (Attending Physician)",
                    "ui_inputs": "• Tabbed or segmented view for:\n  - S (Subjective): Chief complaint, history of present illness\n  - O (Objective): Physical exam findings, systemic observations\n  - A (Assessment): Clinical diagnosis, ICD description\n  - P (Plan): Treatment steps, follow-up timeline\n• Template snippets dropdown (e.g., 'Normal Physical Exam', 'Upper Respiratory URI')",
                    "ui_validation": "Assessment field cannot be empty when finalizing encounter",
                    "ui_processing": "1. Debounced auto-save every 5 seconds to prevent data loss.\n2. Submits payload to PUT /api/v1/encounters/:id/notes.",
                    "ui_outputs": "• Indicator status: 'Draft saved at 10:45 AM'\n• Combined preview panel showing formatted legal clinical note"
                },
                {
                    "code": "UC-FE-03.3",
                    "title": "Encounter Finalization & Clinical Lock",
                    "actor": "Doctor",
                    "ui_inputs": "'Sign & Complete Encounter' button click with confirmation dialog",
                    "ui_validation": "Ensures vitals are entered and SOAP Assessment is completed",
                    "ui_processing": "1. Sends POST /api/v1/encounters/:id/complete.\n2. Transitions encounter status from 'IN_PROGRESS' to 'COMPLETED'.\n3. Converts SOAP inputs from editable forms to read-only clinical archive.",
                    "ui_outputs": "• Badge flips to 'COMPLETED (Signed)' with timestamp and doctor signature stamp\n• Locks note from further inline editing"
                }
            ]
        },
        {
            "id": "MOD-04",
            "title": "CPOE Diagnostic Orders Module (Computerized Physician Order Entry)",
            "desc": "Enables clinicians to place diagnostic laboratory, radiology, and procedural orders, and review fulfillment results.",
            "components": "CpoeOrderCartComponent, OrderCatalogueModalComponent, OrderTrackerListComponent, OrderResultModalComponent",
            "use_cases": [
                {
                    "code": "UC-FE-04.1",
                    "title": "Diagnostic Test Catalogue Search & Order Staging",
                    "actor": "Doctor",
                    "ui_inputs": "• Search order catalogue (Typeahead, e.g. 'CBC', 'Lipid', 'Chest X-Ray', 'ECG')\n• Category filter tabs: [All, Laboratory, Radiology, Cardiology]\n• Priority Toggle: [Routine, Urgent, STAT (High Priority)]\n• Clinical Indication text box (e.g., 'Fever of unknown origin')",
                    "ui_validation": "Clinical indication required for STAT or Imaging orders",
                    "ui_processing": "1. Selected tests add into a temporary 'Order Requisition Cart'.\n2. User can add multiple orders simultaneously (e.g., CBC + Urinalysis + Chest X-Ray).",
                    "ui_outputs": "• Order Cart badge increments count\n• STAT orders highlight in bold red with emergency icon"
                },
                {
                    "code": "UC-FE-04.2",
                    "title": "Batch Digital Sign-Off & Order Transmission",
                    "actor": "Doctor",
                    "ui_inputs": "'Submit & Sign Orders' button click in Order Requisition Cart",
                    "ui_validation": "At least one order in cart",
                    "ui_processing": "1. Sends batch payload to POST /api/v1/orders.\n2. Clears cart on 201 Created and refreshes active order list.",
                    "ui_outputs": "• Success Toast: '3 diagnostic orders submitted successfully'\n• Table updates showing new rows in 'PENDING' status"
                },
                {
                    "code": "UC-FE-04.3",
                    "title": "Diagnostic Results Viewer & Abnormal Value Highlighting",
                    "actor": "Doctor, Nurse, Lab Tech",
                    "ui_inputs": "• Click on order row with status 'COMPLETED'\n• Or 'Enter Results' button (for Lab staff)",
                    "ui_validation": "For result entry: Result narrative or value cannot be blank",
                    "ui_processing": "1. Fetches result CLOB from order details.\n2. Parses text for out-of-range flags (e.g., 'ELEVATED', 'CRITICAL').",
                    "ui_outputs": "• Modal dialog displaying formatted test results, reference ranges, and lab tech sign-off\n• Out-of-range indicators highlighted in red/amber font"
                }
            ]
        },
        {
            "id": "MOD-05",
            "title": "Medication Management & e-Prescribing Module",
            "desc": "Comprehensive medication reconciliation, e-prescribing with automated allergy alert warnings, and active prescription lifecycle controls.",
            "components": "MedicationListComponent, PrescribeModalComponent, AllergyWarningDialogComponent, DiscontinueMedModalComponent",
            "use_cases": [
                {
                    "code": "UC-FE-05.1",
                    "title": "Active Medication Profile & Reconciliation Grid",
                    "actor": "Doctor, Nurse",
                    "ui_inputs": "• Tab filter: [Active Medications, Completed, Discontinued]\n• Search drug name filter",
                    "ui_validation": "None (read-only view)",
                    "ui_processing": "Queries GET /api/v1/prescriptions?patientId={id}&status=ACTIVE.",
                    "ui_outputs": "• Cards / table displaying: Drug Name, Dosage & Unit, Frequency, Route, Prescribing Doctor, Start Date, Days Remaining bar\n• Color status pills: Green for 'ACTIVE', Gray for 'COMPLETED', Red outline for 'DISCONTINUED'"
                },
                {
                    "code": "UC-FE-05.2",
                    "title": "e-Prescribing Drug Form & Frequency Calculator",
                    "actor": "Doctor",
                    "ui_inputs": "• Drug Name (Autocomplete with common pharmaceutical formulas)\n• Dosage (Text, e.g., '500 mg', '10 ml', '20 mcg')\n• Frequency (Dropdown: Once daily, Twice daily [BID], Every 8 hrs [TID], QID, PRN [As needed])\n• Route (Dropdown: Oral, IV, Topical, Inhalation, Subcutaneous)\n• Duration in Days (Number input)\n• Administration Instructions (e.g., 'Take after food with full glass of water')",
                    "ui_validation": "All required fields must be non-empty; duration must be > 0",
                    "ui_processing": "1. Triggers pre-flight allergy check against Patient Chart Header recorded allergies.\n2. If candidate drug matches patient allergy list, intercepts submission.",
                    "ui_outputs": "• Interactive prescription summary preview before final confirmation"
                },
                {
                    "code": "UC-FE-05.3",
                    "title": "Real-Time Drug-Allergy Warning & Physician Override Modal",
                    "actor": "Doctor",
                    "ui_inputs": "• Physician attempts to prescribe a drug that conflicts with patient allergies (e.g., Patient allergic to 'Penicillin', Doctor prescribes 'Amoxicillin')\n• Override Reason (Dropdown: 'Benefit outweighs risk', 'Patient tolerates under supervision', 'Alternative unavailable')\n• Physician PIN / Confirmation checkbox: 'I acknowledge the allergy contraindication'",
                    "ui_validation": "Override reason and explicit checkbox check required to proceed",
                    "ui_processing": "1. Frontend displays high-visibility emergency modal dialog.\n2. If doctor cancels: Reverts back to prescription form without saving.\n3. If doctor submits override: Sends prescription payload with override metadata to POST /api/v1/prescriptions.",
                    "ui_outputs": "• High-contrast Red Alert Modal: 'CRITICAL DRUG ALLERGY WARNING'\n• Displays: 'Patient has documented severe allergy to Penicillin. Amoxicillin is in the penicillin class.'\n• Upon override: Prescription marked with 'ALLERGY OVERRIDE' badge in medication table"
                },
                {
                    "code": "UC-FE-05.4",
                    "title": "Medication Discontinuation Action",
                    "actor": "Doctor",
                    "ui_inputs": "• Click 'Discontinue' button on active medication row\n• Reason for Discontinuation (Dropdown: 'Adverse Reaction', 'Condition Resolved', 'Dosage Adjustment', 'Ineffective')",
                    "ui_validation": "Discontinuation reason required",
                    "ui_processing": "1. Sends PUT /api/v1/prescriptions/:id/discontinue with reason.\n2. Transitions status from 'ACTIVE' to 'DISCONTINUED'.",
                    "ui_outputs": "• Row instantly shifts to 'Discontinued' tab with strikethrough styling\n• Toast: 'Prescription for Amoxicillin has been discontinued'"
                }
            ]
        },
        {
            "id": "MOD-06",
            "title": "Patient Self-Service Portal Module",
            "desc": "A streamlined, user-friendly portal interface designed exclusively for patients to review their own health records securely.",
            "components": "PatientDashboardComponent, MyVitalsTrendComponent, MyMedsCardComponent, MyLabResultsComponent, AfterVisitSummaryPrintComponent",
            "use_cases": [
                {
                    "code": "UC-FE-06.1",
                    "title": "Patient Consolidated Health Dashboard",
                    "actor": "Patient (Self)",
                    "ui_inputs": "Patient logs in; landing page is '/portal'",
                    "ui_validation": "Authentication token verified with ROLE_PATIENT",
                    "ui_processing": "1. PortalService calls GET /api/v1/portal/dashboard.\n2. Aggregates data into 4 clean visual cards: Vitals Overview, Active Prescriptions, Recent Test Results, and Care Team Contacts.",
                    "ui_outputs": "• Welcoming greeting: 'Welcome, John Doe (MRN-2026-0042)'\n• Summary stats: 2 Active Meds | 1 Pending Lab | Next appointment: Oct 12\n• Quick-action buttons: 'View Medications', 'View Lab Results', 'Print Health Summary'"
                },
                {
                    "code": "UC-FE-06.2",
                    "title": "Patient Medication Schedule & Refill Instructions View",
                    "actor": "Patient",
                    "ui_inputs": "Click 'My Medications' tab in portal navigation",
                    "ui_validation": "None",
                    "ui_processing": "Renders active prescriptions with non-technical, patient-friendly formatting.",
                    "ui_outputs": "• Medication cards with pill icons, exact instructions ('Take 1 tablet twice a day with food'), and prescribing doctor name\n• Remaining days counter badge"
                },
                {
                    "code": "UC-FE-06.3",
                    "title": "Patient Diagnostic Results & Lab Reports Viewer",
                    "actor": "Patient",
                    "ui_inputs": "Click 'My Lab Results' tab; click on specific completed test row",
                    "ui_validation": "Only displays results with status == 'COMPLETED'",
                    "ui_processing": "Fetches completed diagnostic reports with plain-language status (e.g. 'Normal', 'Completed').",
                    "ui_outputs": "• Clean expandable card showing test name, date performed, ordering physician, and doctor's remarks\n• 'Download Official Lab Report' button"
                },
                {
                    "code": "UC-FE-06.4",
                    "title": "After-Visit Summary (AVS) Print & Export",
                    "actor": "Patient",
                    "ui_inputs": "Click 'Download / Print Summary' on any completed visit",
                    "ui_validation": "Visit must be in 'COMPLETED' status",
                    "ui_processing": "1. Triggers Angular dedicated print-optimized CSS stylesheet (@media print).\n2. Formats clinical encounter without UI chrome (headers/buttons stripped).",
                    "ui_outputs": "• Native browser print preview modal with cleanly laid-out After-Visit Summary including diagnosis, active prescriptions, and follow-up guidance"
                }
            ]
        },
        {
            "id": "MOD-07",
            "title": "HIPAA Audit Trail & Clinical Administration Module",
            "desc": "Administrative interface for compliance officers and system administrators to inspect security audit logs and user privileges.",
            "components": "AuditLogViewerComponent, AuditFilterBarComponent, UserManagementComponent",
            "use_cases": [
                {
                    "code": "UC-FE-07.1",
                    "title": "HIPAA PHI Access Trail Audit Querying",
                    "actor": "System Administrator, Compliance Officer",
                    "ui_inputs": "• Date Range Picker (Start Date to End Date)\n• Action Filter Dropdown: [All, VIEW_RECORD, CREATE_ORDER, DOCUMENT_SOAP, PRESCRIBE_DRUG, ALLERGY_OVERRIDE]\n• Target Entity Filter: [PATIENT, CLINICAL_NOTE, CPOE_ORDER, PRESCRIPTION]\n• User Search Input (Filter by Staff Username)",
                    "ui_validation": "End Date cannot precede Start Date",
                    "ui_processing": "1. Calls GET /api/v1/audit/logs with pagination.\n2. Renders tabular chronological log with colored action chips.",
                    "ui_outputs": "• Interactive audit table: Timestamp, User, Action, Entity, Entity ID, Client IP\n• Special highlighted rows for ALLERGY_OVERRIDE and security events\n• 'Export to CSV' button"
                }
            ]
        }
    ]

    for mod in modules_data:
        add_heading_styled(doc, f"{mod['id']}: {mod['title']}", level=1)
        
        # Module metadata
        p_desc = doc.add_paragraph()
        run_desc = p_desc.add_run(f"Module Scope: {mod['desc']}\n")
        run_desc.font.italic = True
        run_desc.font.size = Pt(9.5)
        
        run_comp = p_desc.add_run(f"Key Angular Components: {mod['components']}")
        run_comp.font.bold = True
        run_comp.font.color.rgb = RGBColor(27, 85, 155)
        run_comp.font.size = Pt(9.5)
        p_desc.paragraph_format.space_after = Pt(6)

        # Use cases in this module
        for uc in mod["use_cases"]:
            add_heading_styled(doc, f"{uc['code']}: {uc['title']}", level=2)
            uc_table_data = [
                ["Target Actor / Role", uc["actor"]],
                ["UI Inputs & Controls", uc["ui_inputs"]],
                ["Client Validations & Error States", uc["ui_validation"]],
                ["Frontend Processing & Reactive State", uc["ui_processing"]],
                ["Visual Outputs & UI Feedback", uc["ui_outputs"]]
            ]
            create_table_styled(doc, ["Frontend Dimension", "Interaction & Behavior Specification"], uc_table_data, col_widths=[1.9, 4.6])

    # Save Document
    doc.save(output_path)
    print(f"Frontend modules document successfully saved at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "CareConnect_Frontend_Modules_and_UseCases.docx"
    build_frontend_document(out)
